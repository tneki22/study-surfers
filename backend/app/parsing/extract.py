"""Чтение учебных файлов без OCR, БД и внешних сервисов."""

import re
from collections import Counter
from collections.abc import Iterator
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from statistics import median
from uuid import UUID, uuid4
from zipfile import BadZipFile

import pymupdf as fitz
from charset_normalizer import from_bytes
from docx import Document
from docx.text.paragraph import Paragraph
from lxml.etree import XMLSyntaxError

from app.core.errors import DomainError
from app.parsing import docx_content
from app.parsing.archive_limits import check_docx_archive
from app.parsing.quality import MATH_RE, RECOVERY_STEPS, reading_warning, text_warnings
from app.parsing.source_markup import (
    SUPPORTED_KINDS,
    SUPPORTED_LABEL,
    normalize_source,
    source_warnings,
)

NUMBERED_RE = re.compile(r"^\s*(?:\d{1,4}[.)]|(?:Вопрос|Билет|Тема)\s*(?:№\s*)?\d+)[ .:)]*", re.I)
BULLET_RE = re.compile(r"^\s*[•▪◦‣·*+\-–—]\s+")
ORDERED_RE = re.compile(r"^\s*(?:\d{1,4}|[a-zа-я])[.)]\s+")
HEADING_MAX_CHARS = 120
INDENT_STEP = 18
MAX_LIST_LEVEL = 4
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
BULLET_FORMATS = {"bullet", "none"}
LETTER_FORMATS = {
    "lowerLetter": "абвгдежзиклмнопрстуфхцчшщэюя",
    "upperLetter": "АБВГДЕЖЗИКЛМНОПРСТУФХЦЧШЩЭЮЯ",
}


@dataclass(slots=True)
class Block:
    text: str
    element_kind: str = "paragraph"
    id: UUID = field(default_factory=uuid4)
    page: int = 1
    level: int | None = None

    @property
    def kind(self) -> str:
        return self.element_kind


@dataclass(slots=True)
class ExtractedDocument:
    filename: str
    kind: str
    text: str
    blocks: list[Block]
    role: str | None = None
    id: str = field(default_factory=lambda: str(uuid4()))
    warnings: list[dict] = field(default_factory=list)
    sha256: str | None = None
    upload_path: str | None = None

    @property
    def plain(self) -> str:
        return self.text

    @property
    def elements(self) -> list[Block]:
        return self.blocks


def parse_text_page(text: str, page_number: int = 1) -> list[Block]:
    blocks = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            kind, value, level = "heading", heading[2], len(heading[1])
        elif NUMBERED_RE.match(line) and len(line) <= HEADING_MAX_CHARS:
            kind, value, level = "heading", line, 1
        elif BULLET_RE.match(line) or ORDERED_RE.match(line):
            kind, value, level = "list", line, 1
        else:
            kind, value, level = "paragraph", line, None
        blocks.append(Block(value, kind, page=page_number, level=level))
    return blocks


def decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        detected = from_bytes(data).best()
        if detected is not None:
            return str(detected)
        return data.decode("cp1251")


def _pdf_document(filename: str, data: bytes) -> ExtractedDocument:
    pages = []
    warnings = []
    missing_pages, sparse_pages, formula_pages = [], [], []
    with fitz.open(stream=data, filetype="pdf") as document:
        if document.needs_pass:
            raise DomainError(
                422,
                "import_pdf_password",
                "PDF защищён паролем. Сохраните копию без пароля и загрузите её снова.",
                {"filename": filename},
            )
        for page in document:
            lines = []
            has_formulas = False
            for block in page.get_text("dict", sort=True)["blocks"]:
                for line in block.get("lines", []):
                    spans = line["spans"]
                    text = "".join(
                        docx_content.readable(s["text"], s.get("font")) for s in spans
                    ).strip()
                    has_formulas |= bool(MATH_RE.search(text)) or any(
                        s["flags"] & 1
                        or re.search(r"math|symbol|cmsy|cmmi|cmex", s.get("font", ""), re.I)
                        for s in spans
                    )
                    if text:
                        lines.append(
                            (
                                text,
                                max(s["size"] for s in spans),
                                any(s["flags"] & 16 for s in spans),
                            )
                        )
            pages.append(lines)
            length = sum(len(row[0]) for row in lines)
            if not length:
                missing_pages.append(page.number + 1)
            elif length < 20 or any(
                fitz.Rect(image["bbox"]).get_area() > page.rect.get_area() * 0.15
                for image in page.get_image_info()
            ):
                sparse_pages.append(page.number + 1)
            if has_formulas:
                formula_pages.append(page.number + 1)
    for code, message, affected in (
        (
            "pdf_no_text_layer",
            "На некоторых страницах PDF не удалось прочитать текст. "
            "Возможно, это сканы или картинки без распознанного текста.",
            missing_pages,
        ),
        (
            "pdf_text_incomplete",
            "На некоторых страницах PDF прочитано мало текста или есть крупные изображения. "
            "Часть материала может быть сохранена картинками или не распознана. "
            "Проверьте, что текст прочитан полностью.",
            sparse_pages,
        ),
        (
            "pdf_formulas_need_review",
            "В PDF найдены признаки формул. При чтении могли потеряться "
            "дроби, степени, индексы или порядок символов. Сверьте формулы с оригиналом; "
            "если они прочитаны неверно, подготовьте файл заново.",
            formula_pages,
        ),
    ):
        if affected:
            warnings.append(reading_warning(code, message, filename, pages=affected))
    # Считаем повтор по страницам, а не по вхождениям в одной странице.
    counts = Counter(
        text for lines in pages for text in {row[0] for row in lines if len(row[0]) <= 120}
    )
    repeated = {
        text
        for text, count in counts.items()
        if len(pages) >= 2 and count >= max(2, len(pages) / 2)
    }
    sizes = [
        size
        for lines in pages
        for text, size, _ in lines
        if text not in repeated and not text.isdecimal()
    ]
    base_size = median(sizes) if sizes else 12
    blocks = []
    for index, lines in enumerate(pages, 1):
        content = [
            (text, size, bold)
            for text, size, bold in lines
            if text not in repeated and not text.isdecimal()
        ]
        for text, size, bold in content:
            heading = len(text) <= 120 and (size >= base_size * 1.15 or bold)
            if (
                blocks
                and blocks[-1].page == index
                and blocks[-1].element_kind == "paragraph"
                and not heading
                and re.search(r"[а-яё]-$", blocks[-1].text)
                and re.match(r"^[а-яё]", text)
            ):
                blocks[-1].text = blocks[-1].text[:-1] + text
            else:
                blocks.append(Block(text, "heading" if heading else "paragraph", page=index))
    return ExtractedDocument(
        filename, "pdf", "\n".join(b.text for b in blocks), blocks, warnings=warnings
    )


def extract_document(filename: str, data: bytes, role: str | None = None) -> ExtractedDocument:
    kind = Path(filename).suffix.lower().lstrip(".")
    if kind not in SUPPORTED_KINDS:
        hint = " Сохраните документ как DOCX." if kind == "doc" else ""
        raise DomainError(
            415,
            "import_unsupported_type",
            f"Поддерживаются {SUPPORTED_LABEL}." + hint,
            {"filename": filename},
        )
    try:
        if kind == "pdf":
            result = _pdf_document(filename, data)
        elif kind == "docx":
            result = _docx_page(BytesIO(data))
            result.filename = filename
            for warning in result.warnings:
                warning["context"]["filename"] = filename
        else:
            text = decode_text(data)
            warnings = text_warnings(filename, kind, text)
            if kind in {"typ", "tex"}:
                text, dependencies = normalize_source(text, kind)
                warnings.extend(source_warnings(filename, dependencies))
            result = ExtractedDocument(filename, kind, text, parse_text_page(text))
            result.warnings = warnings
    except (
        ValueError,
        UnicodeError,
        fitz.FileDataError,
        KeyError,
        BadZipFile,
        XMLSyntaxError,
    ) as exc:
        raise DomainError(
            422,
            "import_file_invalid",
            f"Не удалось прочитать файл «{filename}». Возможно, файл повреждён "
            "или его содержимое не соответствует расширению. Сохраните исправленную копию. "
            + " ".join(RECOVERY_STEPS),
            reading_warning("import_file_invalid", "", filename)["context"],
        ) from exc
    result.role = role
    if kind in {"pdf", "docx"}:
        result.warnings.extend(text_warnings(filename, kind, result.text))
    return result


def extract_paste(text: str) -> ExtractedDocument:
    return ExtractedDocument(
        "Вставленный текст",
        "paste",
        text,
        parse_text_page(text),
        warnings=text_warnings("Вставленный текст", "paste", text),
    )


def _docx_numbering(document: Document) -> dict[tuple[int, int], tuple[str, str]]:
    """(numId, ilvl) → (формат, шаблон) из numbering.xml.

    Word держит номера пунктов не в тексте абзаца, а в отдельной части документа.
    Без этой карты нумерованный список приезжает без номеров — а в списке
    вопросов номер и есть половина смысла.
    """
    try:
        element = document.part.numbering_part.element
    except (KeyError, AttributeError, NotImplementedError):
        return {}

    abstract: dict[str, dict[int, tuple[str, str]]] = {}
    for node in element.findall(f"{W_NS}abstractNum"):
        levels: dict[int, tuple[str, str]] = {}
        for level in node.findall(f"{W_NS}lvl"):
            fmt = level.find(f"{W_NS}numFmt")
            template = level.find(f"{W_NS}lvlText")
            levels[int(level.get(f"{W_NS}ilvl", "0"))] = (
                fmt.get(f"{W_NS}val", "decimal") if fmt is not None else "decimal",
                template.get(f"{W_NS}val", "%1.") if template is not None else "%1.",
            )
        abstract[node.get(f"{W_NS}abstractNumId", "")] = levels

    mapping: dict[tuple[int, int], tuple[str, str]] = {}
    for node in element.findall(f"{W_NS}num"):
        reference = node.find(f"{W_NS}abstractNumId")
        levels = abstract.get(reference.get(f"{W_NS}val", "")) if reference is not None else None
        if not levels:
            continue
        num_id = int(node.get(f"{W_NS}numId", "0"))
        for ilvl, spec in levels.items():
            mapping[(num_id, ilvl)] = spec
    return mapping


def _numeral(value: int, fmt: str) -> str:
    alphabet = LETTER_FORMATS.get(fmt)
    if alphabet:
        return alphabet[(value - 1) % len(alphabet)]
    return str(value)


def _list_marker(spec: tuple[str, str], ilvl: int, counters: dict[int, int]) -> str:
    fmt, template = spec
    if fmt in BULLET_FORMATS:
        return "—"
    marker = template
    for level, count in counters.items():
        numeral = _numeral(count, fmt if level == ilvl else "decimal")
        marker = marker.replace(f"%{level + 1}", numeral)
    return re.sub(r"%\d", "", marker)


def _docx_num_ref(paragraph: Paragraph) -> tuple[int, int] | None:
    properties = paragraph._p.pPr
    numbering = properties.numPr if properties is not None else None
    if numbering is None or numbering.numId is None:
        return None
    ilvl = numbering.ilvl.val if numbering.ilvl is not None else 0
    return int(numbering.numId.val), int(ilvl or 0)


def _docx_list_level(paragraph: Paragraph) -> int | None:
    """Уровень пункта списка: сначала настоящая нумерация Word, потом отступ слева."""
    properties = paragraph._p.pPr
    numbering = properties.numPr if properties is not None else None
    if numbering is not None:
        raw_level = numbering.ilvl.val if numbering.ilvl is not None else 0
        return 1 + min(MAX_LIST_LEVEL - 1, int(raw_level or 0))
    indent = paragraph.paragraph_format.left_indent
    if indent is None:
        return None
    return 1 + min(MAX_LIST_LEVEL - 1, int(indent.pt // INDENT_STEP))


def _docx_element(
    paragraph: Paragraph, bbox: tuple[float, float, float, float], marker: str, body: str
) -> Block:
    """Элемент абзаца DOCX; `body` — текст с формулами (`docx_content`)."""
    style = paragraph.style.name.lower() if paragraph.style else ""
    if not marker and docx_content.display_only(body):
        return Block(body, "formula")
    text = f"{marker} {body}".strip() if marker else body
    list_level = _docx_list_level(paragraph)
    if marker or "list" in style or BULLET_RE.match(text):
        return Block(text, "list", level=list_level or 1)
    if style.startswith("heading"):
        level_match = re.search(r"(\d+)$", style)
        return Block(text, "heading", level=int(level_match.group(1)) if level_match else 1)
    if NUMBERED_RE.match(text) and len(text) <= HEADING_MAX_CHARS:
        return Block(text, "heading", level=1)
    if ORDERED_RE.match(text):
        return Block(text, "list", level=list_level or 1)
    return Block(text, "paragraph")


def _docx_blocks(document: Document) -> Iterator[Paragraph | object]:
    """Абзацы и таблицы тела документа в их порядке (`document.paragraphs`
    таблиц не видит). Таблица отдаётся своим XML-узлом `w:tbl`."""
    body = document.element.body
    pending = list(body.iterchildren())
    while pending:
        node = pending.pop(0)
        if node.tag == f"{W_NS}p":
            yield Paragraph(node, document._body)
        elif node.tag == f"{W_NS}tbl":
            yield node
        elif node.tag == f"{W_NS}sdt":
            content = node.find(f"{W_NS}sdtContent")
            if content is not None:
                pending[:0] = list(content.iterchildren())


def _docx_page(path) -> ExtractedDocument:
    check_docx_archive(path)
    document = Document(path)
    numbering = _docx_numbering(document)
    counters = {}
    blocks = []
    for block in _docx_blocks(document):
        if not isinstance(block, Paragraph):
            text = docx_content.table_markdown(block)
            if text:
                blocks.append(Block(text, "table"))
            continue
        text = docx_content.paragraph_text(block._p)
        if not text:
            continue
        marker = ""
        reference = _docx_num_ref(block)
        if reference is None and block.style is not None:
            style = block.style
            seen = set()
            while style is not None and style.style_id not in seen:
                seen.add(style.style_id)
                num = style.element.find(f"{W_NS}pPr/{W_NS}numPr")
                if num is not None:
                    num_id, level = num.find(f"{W_NS}numId"), num.find(f"{W_NS}ilvl")
                    if num_id is not None:
                        reference = (
                            int(num_id.get(f"{W_NS}val")),
                            int(level.get(f"{W_NS}val", "0")) if level is not None else 0,
                        )
                        break
                style = style.base_style
        if reference in numbering:
            num_id, ilvl = reference
            counts = counters.setdefault(num_id, {})
            counts[ilvl] = counts.get(ilvl, 0) + 1
            for deeper in list(counts):
                if deeper > ilvl:
                    del counts[deeper]
            marker = _list_marker(numbering[reference], ilvl, counts)
        blocks.append(_docx_element(block, (0, 0, 1, 1), marker, text))
    result = ExtractedDocument("Документ DOCX", "docx", "\n".join(b.text for b in blocks), blocks)
    if result.text and any(
        node.tag.rsplit("}", 1)[-1] in {"drawing", "pict", "object"}
        for node in document.element.body.iter()
        if isinstance(node.tag, str)
    ):
        result.warnings.append(
            reading_warning(
                "document_images_skipped",
                "В документе есть картинки или встроенные объекты. "
                "Текст и формулы внутри них не прочитаны и могут не попасть в карточки.",
                result.filename,
            )
        )
    return result
