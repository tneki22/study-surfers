import io
import json

import pymupdf as fitz
import pytest
from docx import Document
from docx.oxml import parse_xml
from test_ai_client import configured as configured_model
from test_jobs import completion, inject, wait_job, wire

from app.ai.provider import FakeTransport, ProviderCompletion, ProviderUsage
from app.parsing.extract import extract_document, extract_paste
from app.parsing.quality import broken_math, unreadable_text

configured = configured_model


def test_scan_pages_are_grouped_and_offer_both_recovery_options():
    with fitz.open() as pdf:
        pdf.new_page()
        pdf.new_page()
        result = extract_document("scan.pdf", pdf.tobytes())
    assert len(result.warnings) == 1
    warning = result.warnings[0]
    assert warning["context"]["pages"] == [1, 2]
    assert len(warning["context"]["recovery_steps"]) == 2
    assert "OCR" in warning["context"]["recovery_steps"][0]
    assert "другую модель" in warning["context"]["recovery_steps"][1]
    assert "LaTeX" in warning["context"]["formula_help"]


def test_mixed_pdf_keeps_readable_material_and_identifies_missing_page():
    with fitz.open() as pdf:
        pdf.new_page().insert_text((30, 30), "Material about physics and energy conservation.")
        pdf.new_page()
        result = extract_document("mixed.pdf", pdf.tobytes())
    assert "energy conservation" in result.text
    assert result.warnings[0]["context"]["pages"] == [2]


def test_pdf_sparse_text_and_formula_layout_warning():
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((30, 30), "E = mc")
        page.insert_text((70, 25), "2", fontsize=7)
        result = extract_document("formula.pdf", pdf.tobytes())
    assert "pdf_text_incomplete" in {w["code"] for w in result.warnings}
    assert "pdf_no_text_layer" not in {w["code"] for w in result.warnings}


def test_pdf_detects_math_font_and_warns_about_structure():
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((30, 30), "A sufficiently long paragraph about mathematical limits.")
        page.insert_text((30, 60), "a", fontname="symb")
        result = extract_document("math.pdf", pdf.tobytes())
    assert "pdf_formulas_need_review" in {w["code"] for w in result.warnings}


def test_docx_image_only_and_embedded_object_are_explained():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph._p.append(
        parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:object/></w:r>"
        )
    )
    stream = io.BytesIO()
    document.save(stream)
    assert (
        extract_document("scan.docx", stream.getvalue()).warnings[0]["code"] == "document_no_text"
    )
    paragraph.add_run("Материал о формулах и физических величинах")
    stream = io.BytesIO()
    document.save(stream)
    warning = extract_document("notes.docx", stream.getvalue()).warnings[0]
    assert warning["code"] == "document_images_skipped"
    assert warning["context"]["filename"] == "notes.docx"


@pytest.mark.parametrize(
    "text",
    ["Текст \ufffd\ufffd испорчен", "\ue000\ue001 текст", "Привет".encode().decode("cp1251")],
)
def test_garbled_text_has_recovery_warning(text):
    assert unreadable_text(text)
    assert extract_paste(text).warnings[0]["code"] == "document_text_unreadable"


@pytest.mark.parametrize(
    "text",
    ["Обычный русский текст с α, β и ∑", "Знак \ufffd обсуждается в тексте", "100 долларов: 100$"],
)
def test_readable_text_is_not_garbled(text):
    assert not unreadable_text(text)


@pytest.mark.parametrize(
    "text", [r"Формула $x^{2$", r"Формула $x+1", r"Формула \[x+1", r"Формула $$x+1$"]
)
def test_invalid_formula_delimiters_are_explained(text):
    assert extract_document("notes.md", text.encode()).warnings[0]["code"] == "formulas_invalid"


@pytest.mark.parametrize(
    "text",
    [
        r"Стоимость 100$",
        r"Интервал $[a,b)$",
        r"Доля $\frac{a}{b}$",
        r"Пример: `$x+1`",
        "```tex\n$x+1\n```",
    ],
)
def test_valid_math_code_and_currency_are_not_rejected(text):
    assert not broken_math(text, "md")


@pytest.mark.parametrize(
    "filename,source,formula",
    [
        ("notes.TYP", "= Дроби\nФормула: $ frac(a, b) + x_(i+1) $.", "$ frac(a, b) + x_(i+1) $"),
        (
            "notes.tex",
            r"\section{Дроби}" + "\n" + r"Формула: $\frac{a}{b} + x_{i+1}$.",
            r"$\frac{a}{b} + x_{i+1}$",
        ),
    ],
)
def test_source_import_keeps_formulas_and_headings(client, filename, source, formula):
    document = extract_document(filename, source.encode())
    assert document.blocks[0].kind == "heading"
    assert document.blocks[0].text == "Дроби"
    assert formula in document.text
    assert document.warnings == []
    response = client.post(
        "/api/imports/quick?dry_run=true",
        files={"files[]": (filename, source.encode())},
        data={"options": json.dumps({"roles": {filename: "notes"}})},
    )
    assert response.status_code == 200, response.text
    assert formula in response.json()["questions"][0]["answer_text"]


def test_source_dependencies_are_not_silently_ignored():
    result = extract_document("main.typ", b'#include "chapter.typ"\n= Topic\nText')
    assert result.warnings[0]["code"] == "source_requires_review"
    assert "chapter.typ" in result.text


def test_typst_urls_strings_and_escaped_dollars_are_preserved():
    source = '= Ссылки\nhttps://typst.app\n#let label = "$x+1$"\nЦена: \\$100\n$frac(a,b)$'
    result = extract_document("source.typ", source.encode())
    assert "https://typst.app" in result.text
    assert '"$x+1$"' in result.text
    assert "Цена: \\$100" in result.text
    assert "```typst\n$frac(a,b)$\n```" in result.text


def test_invalid_source_formula_is_checked_before_typst_tagging():
    result = extract_document("formula.typ", b"= Math\n$ frac(a, b) + x_(i+1)")
    assert result.warnings[0]["code"] == "formulas_invalid"


def test_latex_math_environment_must_be_closed():
    assert broken_math(r"\begin{equation}x=1", "tex")
    assert not broken_math(r"\begin{align*}x=1\end{align*}", "tex")


def test_single_broken_symbol_in_formula_is_reported():
    result = extract_document("notes.md", "Формула $x + \ufffd$".encode())
    assert result.warnings[0]["code"] == "formulas_invalid"


def test_scan_job_persists_reading_help_without_calling_model(configured, client, monkeypatch):
    transport = FakeTransport()
    inject(monkeypatch, transport)
    with fitz.open() as pdf:
        pdf.new_page()
        data = pdf.tobytes()
    response = client.post("/api/imports/quick", files={"files[]": ("scan.pdf", data)})
    assert response.status_code == 202
    result = wait_job(client, response.json()["job_id"])
    assert result["error"]["code"] == "import_no_questions"
    assert result["warnings"][0]["context"]["recovery_steps"]
    source = client.get(f"/api/decks/{result['deck_id']}").json()["sources"][0]
    assert source["warnings"] == result["warnings"]
    assert transport.complete_calls == 0


def test_typst_job_sends_original_formula_and_conversion_instruction(
    configured, client, monkeypatch
):
    outline = {
        "questions": [
            {
                "start": 0,
                "question": "Энергия",
                "section": "",
                "key_points": ["Энергия покоя тела E = mc^2"],
            }
        ],
        "skip": [],
    }
    transport = FakeTransport(
        completions=[
            ProviderCompletion(json.dumps(outline), "test-model", ProviderUsage(10, 20)),
            completion([wire(explanation=r"Формула $E=mc^2$.")]),
        ]
    )
    inject(monkeypatch, transport)
    source = "= Энергия\nМатериал об энергии. Формула $ E = m c^2 $."
    response = client.post(
        "/api/imports/quick",
        files={"files[]": ("physics.typ", source.encode())},
        data={"options": json.dumps({"roles": {"physics.typ": "notes"}, "items_per_question": 1})},
    )
    assert response.status_code == 202
    result = wait_job(client, response.json()["job_id"])
    assert result["state"] == "done", result
    # Сначала модель размечает вопросы по материалу, затем генерирует карточки.
    assert "Энергия" in transport.complete_requests[0]["messages"][1]["content"]
    request = transport.complete_requests[1]
    assert "равнозначный LaTeX" in request["messages"][0]["content"]
    # Конспект: модель знает происхождение материала и главное по теме из программы.
    content = request["messages"][1]["content"]
    context = json.loads(content.removeprefix("<context>").removesuffix("</context>"))
    assert context["material"] == "notes"
    assert context["key_points"] == ["Энергия покоя тела E = mc^2"]
    assert "neighbors" not in context
    assert "$ E = m c^2 $" in request["messages"][1]["content"]
    items = client.get(f"/api/decks/{result['deck_id']}/items").json()["items"]
    assert "$E=mc^2$" in items[0]["explanation"]
    deck = client.get(f"/api/decks/{result['deck_id']}")
    assert deck.status_code == 200, deck.text
    assert deck.json()["sources"][0]["kind"] == "typ"


def test_corrupt_and_password_pdf_messages(client):
    response = client.post(
        "/api/imports/quick?dry_run=true", files={"files[]": ("broken.pdf", b"invalid")}
    )
    assert response.status_code == 422
    assert "OCR" in response.json()["detail"]
    with fitz.open() as pdf:
        pdf.new_page()
        data = pdf.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="test", owner_pw="owner")
    response = client.post(
        "/api/imports/quick?dry_run=true", files={"files[]": ("locked.pdf", data)}
    )
    assert response.status_code == 422
    assert response.json()["code"] == "import_pdf_password"
    assert "без пароля" in response.json()["detail"]
