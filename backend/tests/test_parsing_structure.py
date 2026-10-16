import io
import json
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import pymupdf as fitz
import pytest
from docx import Document
from docx.oxml import parse_xml

from app.core.errors import DomainError
from app.parsing.classify import classify_documents
from app.parsing.extract import extract_document, extract_paste
from app.parsing.prompts import OUTLINE_SYSTEM_PROMPT, outline_prompt
from app.pipeline.structure import StructuredQuestion, _ai_sections, structure_documents

Q_TEXT = "1. Транзакции и блокировки\n2. Использование индексов\n3. Реляционные базы данных"


def test_cp1251_and_utf16_are_read_correctly():
    for encoding in ["cp1251", "utf-16", "utf-8-sig"]:
        document = extract_document("questions.txt", Q_TEXT.encode(encoding))
        assert document.text == Q_TEXT
        assert len(document.blocks) == 3


def test_docx_auto_numbering_inherited_from_style():
    document = Document()
    for text in ["Транзакции", "Индексы", "Базы данных"]:
        document.add_paragraph(text, style="List Number")
    stream = io.BytesIO()
    document.save(stream)
    parsed = extract_document("questions.docx", stream.getvalue())
    assert parsed.text.splitlines() == ["1. Транзакции", "2. Индексы", "3. Базы данных"]


def test_docx_explicit_auto_numbering():
    document = Document()
    for text in ["Первый вопрос", "Второй вопрос", "Третий вопрос"]:
        paragraph = document.add_paragraph(text)
        paragraph._p.get_or_add_pPr().append(
            parse_xml(
                '<w:numPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:ilvl w:val="0"/><w:numId w:val="5"/></w:numPr>'
            )
        )
    stream = io.BytesIO()
    document.save(stream)
    parsed = extract_document("questions.docx", stream.getvalue())
    assert parsed.text.splitlines() == ["1. Первый вопрос", "2. Второй вопрос", "3. Третий вопрос"]


def test_pdf_headings_without_headers_and_lowercase_hyphen():
    pdf = fitz.open()
    for number in range(3):
        page = pdf.new_page()
        page.insert_text((35, 30), "Repeated header", fontsize=9)
        page.insert_text((35, 90), f"Heading {number}", fontsize=18)
        page.insert_text(
            (35, 130),
            f"Paragraph {number}: sufficiently long text about databases and transactions.",
            fontsize=12,
        )
        page.insert_text((35, 180), str(number + 1), fontsize=9)
    parsed = extract_document("answers.pdf", pdf.tobytes())
    assert "Repeated header" not in parsed.text
    assert [b.text for b in parsed.blocks if b.kind == "heading"] == [
        "Heading 0",
        "Heading 1",
        "Heading 2",
    ]
    assert not any(b.text.isdecimal() for b in parsed.blocks)
    assert parsed.warnings == []


def test_pdf_scan_warning_and_doc_hint():
    pdf = fitz.open()
    pdf.new_page()
    parsed = extract_document("scan.pdf", pdf.tobytes())
    assert parsed.warnings[0]["code"] == "pdf_no_text_layer"
    assert parsed.warnings[0]["context"]["pages"] == [1]
    with pytest.raises(DomainError) as error:
        extract_document("legacy.doc", b"anything")
    assert error.value.code == "import_unsupported_type"
    assert "DOCX" in error.value.detail


async def test_questions_without_answers():
    result = await structure_documents([extract_document("questions.txt", Q_TEXT.encode())])
    assert len(result) == 3
    assert all(q.answer_text is None and q.match_method == "none" for q in result)


async def test_russian_question_markers_and_letter_lists():
    for text in [
        "Вопрос 1. Транзакции\nВопрос 2. Индексы\nВопрос 3. Таблицы",
        "а) Транзакции\nб) Индексы\nв) Таблицы",
    ]:
        document = extract_paste(text)
        document.role = "questions"
        result = await structure_documents([document])
        assert [q.text for q in result] == ["Транзакции", "Индексы", "Таблицы"]


async def test_exact_fuzzy_and_orphan_answer_sections():
    text = (
        "# Транзакции и блокировки\nОтвет: Это точный ответ.\n# Использование об индексах\nОтвет: Это нечёткий ответ.\n# Другой раздел\n"
        + "Потерять этот материал нельзя. " * 20
    )
    # Проверенная Tentex пара со служебным «о/об».
    questions = extract_document(
        "questions.txt",
        "1. Транзакции и блокировки\n2. Использование о индексах\n3. Реляционные базы данных".encode(),
    )
    answers = extract_document("answers.md", text.encode(), role="answers")
    result = await structure_documents([questions, answers])
    assert result[0].answer_text == "Это точный ответ."
    assert result[0].match_method == "exact"
    assert result[1].answer_text == "Это нечёткий ответ."
    assert result[1].match_method == "fuzzy"
    assert result[2].answer_text is None
    assert result[3].synthetic
    assert result[3].text == "Другой раздел"


async def test_numbered_order_matching():
    answers = "1. Совсем другое название\nПервый ответ.\n2. Вторая тема\nВторой ответ.\n3. Третья тема\nТретий ответ."
    result = await structure_documents(
        [
            extract_document("questions.txt", Q_TEXT.encode()),
            extract_document("answers.txt", answers.encode(), role="answers"),
        ]
    )
    assert [q.match_method for q in result] == ["numbered_order"] * 3
    assert [q.answer_text for q in result] == ["Первый ответ.", "Второй ответ.", "Третий ответ."]


async def test_combined_and_notes_keep_verbatim_bodies():
    text = (
        "# Транзакции\n"
        + "Транзакция атомарна. " * 12
        + "\n# Индексы\n"
        + "Индекс ускоряет поиск. " * 12
    )
    combined = extract_document("combined.md", text.encode())
    await classify_documents([combined])
    assert combined.role == "combined"
    result = await structure_documents([combined])
    assert len(result) == 2
    assert result[0].text == "Транзакции"
    assert result[0].answer_text == "Транзакция атомарна. " * 11 + "Транзакция атомарна."
    assert all(q.match_method == "combined" and not q.synthetic for q in result)
    notes = extract_paste("Материал конспекта. " * 400)
    notes.role = "notes"
    result = await structure_documents([notes])
    assert len(result) >= 3
    assert all(q.synthetic and len(q.answer_text) <= 2500 for q in result)


class FakeClient:
    def __init__(self, payloads):
        self.payloads = list(payloads)
        self.calls = []

    async def complete(self, **kwargs):
        self.calls.append(kwargs)
        return kwargs["response_model"].model_validate(self.payloads.pop(0))


async def test_dirty_list_through_real_ai_client_and_fake_transport(client):
    from app.ai.client import AiClient
    from app.ai.models import AiModelCatalogEntry, AiProviderConnection, AiSettings
    from app.ai.provider import FakeTransport, ProviderCompletion, ProviderUsage
    from app.db import session_factory

    provider_id = str(uuid4())
    now = datetime.now(timezone.utc)
    with session_factory()() as session:
        session.add(
            AiProviderConnection(
                id=provider_id, label="Тестовый локальный", base_url="http://localhost:1234/v1"
            )
        )
        session.flush()
        session.add(
            AiModelCatalogEntry(
                provider_id=provider_id,
                model_id="test",
                display_name="Тестовая модель",
                input_modalities=["text"],
                output_modalities=["text"],
                pricing_snapshot_at=now,
                catalog_snapshot_at=now,
            )
        )
        settings = session.get(AiSettings, 1)
        if settings is None:
            settings = AiSettings(id=1)
            session.add(settings)
        settings.default_text_provider_id, settings.default_text_model_id = provider_id, "test"
        settings.external_models_enabled = True
        session.commit()
    payload = {
        "items": [
            {
                "kind": "question",
                "title": "Что такое транзакция?",
                "subpoints": [],
                "source_indices": [1, 2],
                "ticket_index": 0,
            }
        ],
        "dropped": [],
        "changes": [],
        "warnings": [],
    }
    transport = FakeTransport(
        completions=[
            ProviderCompletion(
                json.dumps(payload, ensure_ascii=False), "test", ProviderUsage(12, 10)
            )
        ]
    )
    document = extract_paste("Транзак ции\nопреде ление и свойства")
    document.role = "questions"
    result = await structure_documents(
        [document], AiClient(transport=transport, retry_backoff=(0,))
    )
    assert result[0].text == "Что такое транзакция?"
    assert transport.complete_calls == 1


async def test_b4_fixture_documents_extract_and_match():
    fixtures = Path(__file__).parent / "fixtures"
    questions = extract_document("questions.txt", (fixtures / "questions_cp1251.txt").read_bytes())
    answers = extract_document(
        "answers.docx", (fixtures / "answers.docx").read_bytes(), role="answers"
    )
    result = await structure_documents([questions, answers])
    assert len(result) == 3
    assert all(q.match_method == "exact" and q.answer_text for q in result)
    pdf = extract_document("text.pdf", (fixtures / "text.pdf").read_bytes())
    assert len([b for b in pdf.blocks if b.kind == "heading"]) == 3
    combined = await structure_documents(
        [extract_document("combined.md", (fixtures / "combined.md").read_bytes())]
    )
    assert len(combined) == 2


async def test_dirty_list_uses_ai_repair_with_source_indices():
    client = FakeClient(
        [
            {
                "items": [
                    {
                        "kind": "question",
                        "title": "Что такое транзакция?",
                        "subpoints": [],
                        "source_indices": [1, 2],
                        "ticket_index": 0,
                    }
                ],
                "dropped": [],
                "changes": [],
                "warnings": [],
            }
        ]
    )
    dirty = extract_paste("Транзак ции\nопреде ление и свойства")
    dirty.role = "questions"
    result = await structure_documents([dirty], client)
    assert result[0].text == "Что такое транзакция?"
    assert client.calls[0]["role"] == "extract"


async def test_ai_marks_boundaries_server_slices_original_answers():
    answers = extract_document(
        "answers.md",
        "# Незнакомый раздел\nДословный ответ с формулой $x^2$.\n# Другой раздел\nСледующий дословный ответ.".encode(),
        role="answers",
    )
    client = FakeClient(
        [
            {
                "boundaries": [
                    {"candidate": 0, "question": 2, "confidence": "low", "note": ""},
                    {"candidate": 1, "question": 1, "confidence": "low", "note": ""},
                ],
                "skipped_questions": [3],
            }
        ]
    )
    result = await structure_documents(
        [extract_document("questions.txt", Q_TEXT.encode()), answers], client
    )
    assert result[0].answer_text == "Следующий дословный ответ."
    assert result[1].answer_text == "Дословный ответ с формулой $x^2$."
    assert result[1].match_method == "ai_section"
    assert result[2].answer_text is None
    assert "Дословный ответ с формулой" not in client.calls[0]["messages"][1]["content"]


async def test_user_roles_override_classification():
    document = extract_document("questions.txt", Q_TEXT.encode())
    await classify_documents([document], roles={"questions.txt": "notes"})
    assert document.role == "notes"


async def test_classification_batches_large_document_sets():
    documents = [
        extract_document(f"{i}.txt", ("Учебный материал " * 1000).encode()) for i in range(8)
    ]
    client = FakeClient([{"roles": ["notes"] * count} for count in [3, 3, 2]])
    await classify_documents(documents, ai_client=client)
    assert all(d.role == "notes" for d in documents)
    assert len(client.calls) == 3
    for call in client.calls:
        payload = json.loads(call["messages"][1]["content"])
        assert len(payload) <= 3 and all(len(d["text"]) <= 2000 for d in payload)


async def test_local_answer_batches_keep_global_question_numbers():
    document = extract_document(
        "a.md", "# Неизвестный раздел\nДословный ответ.".encode(), role="answers"
    )
    questions = [
        StructuredQuestion(i, str(i), f"Тема {i} " + "подробности " * 100) for i in range(1, 26)
    ]
    client = FakeClient(
        [
            {
                "boundaries": [{"candidate": 0, "question": 0, "confidence": "high", "note": ""}],
                "skipped_questions": [],
            },
            {
                "boundaries": [{"candidate": 0, "question": 25, "confidence": "low", "note": ""}],
                "skipped_questions": [],
            },
        ]
    )
    client.resolve = lambda role: SimpleNamespace(
        provider=SimpleNamespace(base_url="http://localhost:1234/v1")
    )
    sections = await _ai_sections(document, questions, client)
    assert len(sections) == 1 and sections[0].question == 25
    assert sections[0].body == "Дословный ответ."
    assert len(client.calls) == 2
    for call in client.calls:
        payload = json.loads(call["messages"][1]["content"])
        assert len(payload["questions"]) <= 24
        assert all(len(q) <= 140 for q in payload["questions"].values())
    assert list(json.loads(client.calls[1]["messages"][1]["content"])["questions"]) == ["25"]


async def test_ai_outline_decides_questions_in_unstructured_source():
    text = "\n".join(
        [
            "#set page(margin: 1cm)",
            "= Раскраска графов",
            "definition[",
            "Хроматическое число — наименьшее число цветов правильной раскраски.",
            "]",
            "Пусть",
            "$chi(G) <= Delta(G) + 1$",
            "= Клики",
            "definition[",
            "Клика — полный подграф графа.",
            "]",
        ]
    )
    document = extract_document("main.typ", text.encode())
    document.role = "notes"
    client = FakeClient(
        [
            {
                "questions": [
                    {
                        "start": 1,
                        "question": "Раскраска графов и хроматическое число",
                        "section": "Графы",
                        "key_points": ["Хроматическое число — минимум цветов", " ", "x" * 300],
                    },
                    {"start": 99, "question": "Строка вне части", "section": "", "key_points": []},
                    {"start": 7, "question": "Клики в графе", "section": "", "key_points": []},
                ],
                "skip": [{"start": 0, "end": 0}],
            }
        ]
    )
    # Код нашёл бы «вопрос» в каждом заголовке и пункте; число вопросов решает модель.
    result = await structure_documents([document], client)
    assert [q.text for q in result] == ["Раскраска графов и хроматическое число", "Клики в графе"]
    assert [q.section for q in result] == ["Графы", "Графы"]
    assert "#set page" not in result[0].answer_text
    assert "Хроматическое число" in result[0].answer_text
    assert "Клика — полный подграф" in result[1].answer_text
    assert all(q.match_method == "ai_outline" and q.synthetic for q in result)
    assert result[0].key_points == ["Хроматическое число — минимум цветов", "x" * 160]
    assert result[1].key_points == []
    payload = json.loads(client.calls[0]["messages"][1]["content"])
    assert payload["part"] == 1 and payload["parts"] == 1


async def test_ai_outline_follows_learner_wishes():
    document = extract_document("lab.txt", "Цель работы\nКонтрольные вопросы\n".encode())
    document.role = "notes"
    client = FakeClient([{"questions": [], "skip": []}, {"questions": [], "skip": []}])
    wish = "Это лабораторная: карточки только по контрольным вопросам в конце."
    await structure_documents([document], client, learner_prompt=wish)
    system = client.calls[0]["messages"][0]["content"]
    assert system.startswith(OUTLINE_SYSTEM_PROMPT)
    assert system.endswith(f"<learner>\n{wish}\n</learner>")
    assert outline_prompt("   ") == OUTLINE_SYSTEM_PROMPT


def topic(start, question, section="", key_points=()):
    return {"start": start, "question": question, "section": section, "key_points": [*key_points]}


def long_notes():
    document = extract_paste(
        "\n".join(f"Строка материала {i}: " + "текст " * 60 for i in range(400))
    )
    document.role = "notes"
    from app.pipeline.structure import _outline_parts

    parts = _outline_parts(document.blocks, 30000)
    return document, parts, parts[1][0][0]


async def test_ai_outline_parts_keep_global_lines_and_drop_seam_duplicate():
    lines = [f"Строка материала {i}: " + "текст " * 60 for i in range(400)]
    document = extract_paste("\n".join(lines))
    document.role = "combined"
    from app.pipeline.structure import _outline_parts

    parts = _outline_parts(document.blocks, 30000)
    assert len(parts) >= 2
    seam = parts[1][0][0]
    first = {"questions": [topic(0, "Первая тема")], "skip": []}
    # Вторая часть начинается продолжением той же темы и повторяет её.
    repeat = {"questions": [topic(seam, "Первая тема")], "skip": []}
    rest = [{"questions": [], "skip": []}] * (len(parts) - 2)
    client = FakeClient([first, repeat, *rest])
    result = await structure_documents([document], client)
    assert len(result) == 1
    assert f"Строка материала {seam}" in result[0].answer_text
    assert not result[0].synthetic


async def test_program_review_merges_seam_topics_and_aligns_sections():
    document, parts, seam = long_notes()
    first = {"questions": [topic(0, "Графы", "Модуль 1", ["Определение графа"])], "skip": []}
    # На стыке вторая часть назвала продолжение темы иначе — код такой дубль не узнает.
    second = {
        "questions": [
            topic(seam, "Связность графа", "", ["Компонента связности"]),
            topic(seam + 5, "Деревья", "Модуль I", ["Определение дерева"]),
        ],
        "skip": [],
    }
    rest = [{"questions": [], "skip": []}] * (len(parts) - 2)
    program = {
        "topics": [
            {"start": 0, "merge": False, "question": "Графы", "section": "Модуль 1"},
            {"start": seam, "merge": True, "question": "Графы и связность", "section": ""},
            {"start": seam + 5, "merge": False, "question": "Деревья", "section": "Модуль 1"},
        ]
    }
    client = FakeClient([first, second, *rest, program])
    result = await structure_documents([document], client)
    assert [q.text for q in result] == ["Графы и связность", "Деревья"]
    assert [q.section for q in result] == ["Модуль 1", "Модуль 1"]
    assert result[0].key_points == ["Определение графа", "Компонента связности"]
    assert f"Строка материала {seam + 1}" in result[0].answer_text
    payload = json.loads(client.calls[-1]["messages"][1]["content"])
    assert [t["start"] for t in payload["topics"]] == [0, seam, seam + 5]
    assert payload["topics"][1]["key_points"] == ["Компонента связности"]


async def test_program_review_that_loses_topics_is_ignored():
    document, parts, seam = long_notes()
    first = {"questions": [topic(0, "Графы")], "skip": []}
    second = {"questions": [topic(seam, "Деревья")], "skip": []}
    rest = [{"questions": [], "skip": []}] * (len(parts) - 2)
    lost = {"topics": [{"start": 0, "merge": False, "question": "Всё о графах", "section": ""}]}
    client = FakeClient([first, second, *rest, lost])
    result = await structure_documents([document], client)
    assert [q.text for q in result] == ["Графы", "Деревья"]
